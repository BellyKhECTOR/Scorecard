"""PowerPoint extractor."""

import io

from pptx import Presentation

from src.extractors.base import ChunkDraft, ExtractionResult


class PowerPointExtractor:
    method = "pptx"

    def extract(self, file_data: bytes, extension: str = ".pptx") -> ExtractionResult:
        result = ExtractionResult(method="pptx")
        prs = Presentation(io.BytesIO(file_data))
        all_text: list[str] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_parts: list[str] = []
            title = None
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
                if title:
                    slide_parts.append(title)

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    if shape != slide.shapes.title:
                        slide_parts.append(shape.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            slide_parts.append(" | ".join(cells))

            notes_slide = slide.notes_slide
            if notes_slide and notes_slide.notes_text_frame:
                notes = notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"[Notes] {notes}")

            if slide_parts:
                content = "\n".join(slide_parts)
                all_text.append(content)
                result.chunks.append(ChunkDraft(
                    chunk_index=slide_num - 1,
                    content=content,
                    page_number=slide_num,
                    section_heading=title,
                    content_type="text",
                ))

        result.full_text = "\n\n".join(all_text)
        result.status = "completed" if result.full_text else "partial"
        return result
