"""Excel spreadsheet extractor for DDQ questionnaires."""

import io
import logging
from typing import Any

import openpyxl
import xlrd

from src.extractors.base import ChunkDraft, ExtractionResult, QAPairDraft
from src.extractors.excel_utils import (
    detect_header_row,
    find_answer_column,
    find_question_column,
    is_likely_table_not_qa,
    is_section_heading,
    normalise_cell_value,
    normalise_header,
)

logger = logging.getLogger(__name__)


class ExcelExtractor:
  method = "excel"

  def extract(self, file_data: bytes, extension: str) -> ExtractionResult:
    if extension == ".csv":
      return self._extract_csv(file_data)
    if extension == ".xls":
      return self._extract_xls(file_data)
    return self._extract_xlsx(file_data)

  def _extract_xlsx(self, file_data: bytes) -> ExtractionResult:
    result = ExtractionResult(method="excel_xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
    all_text_parts: list[str] = []

    for sheet_name in wb.sheetnames:
      ws = wb[sheet_name]
      if ws.sheet_state == "hidden":
        result.warnings.append(f"Skipped hidden sheet: {sheet_name}")
        logger.info("Skipped hidden sheet: %s", sheet_name)
        continue

      rows = list(ws.iter_rows(values_only=True))
      if not rows:
        continue

      merged_map = self._build_merged_map_xlsx(ws)
      sheet_result = self._process_sheet(rows, sheet_name, merged_map)
      result.qa_pairs.extend(sheet_result.qa_pairs)
      result.chunks.extend(sheet_result.chunks)
      result.warnings.extend(sheet_result.warnings)
      if sheet_result.full_text:
        all_text_parts.append(sheet_result.full_text)

    result.full_text = "\n\n".join(all_text_parts)
    result.status = "completed" if result.qa_pairs or result.chunks else "partial"
    return result

  def _extract_xls(self, file_data: bytes) -> ExtractionResult:
    result = ExtractionResult(method="excel_xls")
    wb = xlrd.open_workbook(file_contents=file_data)
    all_text_parts: list[str] = []

    for sheet_idx in range(wb.nsheets):
      ws = wb.sheet_by_index(sheet_idx)
      if ws.visibility != 0:
        result.warnings.append(f"Skipped hidden sheet: {ws.name}")
        continue
      rows = [tuple(ws.row_values(r)) for r in range(ws.nrows)]
      sheet_result = self._process_sheet(rows, ws.name, {})
      result.qa_pairs.extend(sheet_result.qa_pairs)
      result.chunks.extend(sheet_result.chunks)
      result.warnings.extend(sheet_result.warnings)
      if sheet_result.full_text:
        all_text_parts.append(sheet_result.full_text)

    result.full_text = "\n\n".join(all_text_parts)
    result.status = "completed" if result.qa_pairs or result.chunks else "partial"
    return result

  def _extract_csv(self, file_data: bytes) -> ExtractionResult:
    import csv
    text = file_data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [tuple(row) for row in reader]
    result = self._process_sheet(rows, "CSV", {})
    result.method = "csv"
    result.full_text = "\n".join(
      f"{q.question}\t{q.answer}" for q in result.qa_pairs
    )
    result.status = "completed" if result.qa_pairs else "partial"
    return result

  def _build_merged_map_xlsx(self, ws) -> dict[tuple[int, int], Any]:
    merged_map: dict[tuple[int, int], Any] = {}
    for merged_range in ws.merged_cells.ranges:
      min_row, min_col = merged_range.min_row, merged_range.min_col
      top_left = ws.cell(min_row, min_col).value
      for row in range(merged_range.min_row, merged_range.max_row + 1):
        for col in range(merged_range.min_col, merged_range.max_col + 1):
          merged_map[(row, col)] = top_left
    return merged_map

  def _get_cell(self, row: tuple, row_num: int, col: int, merged_map: dict) -> str:
    if merged_map:
      val = merged_map.get((row_num, col + 1))
      if val is not None:
        return normalise_cell_value(val)
    if col < len(row):
      return normalise_cell_value(row[col])
    return ""

  def _process_sheet(
    self, rows: list[tuple], sheet_name: str, merged_map: dict
  ) -> ExtractionResult:
    result = ExtractionResult(method="excel_sheet")
    if not rows:
      return result

    header_row_idx, headers = detect_header_row(rows)
    section_heading: str | None = None
    text_lines: list[str] = []

    if header_row_idx is not None:
      q_col = find_question_column(headers)
      a_col = find_answer_column(headers, q_col)
      if q_col is None or a_col is None:
        result.warnings.append(f"Sheet '{sheet_name}': incomplete header detection")
        return self._two_column_fallback(rows, sheet_name, merged_map)

      data_rows = rows[header_row_idx + 1:]
      if is_likely_table_not_qa(data_rows, q_col, a_col):
        result.warnings.append(f"Sheet '{sheet_name}': detected as table, skipping Q&A extraction")
        chunk_text = "\n".join(
          " | ".join(normalise_cell_value(c) for c in row if c) for row in data_rows[:50]
        )
        if chunk_text.strip():
          result.chunks.append(ChunkDraft(
            chunk_index=0, content=chunk_text, sheet_name=sheet_name, content_type="table"
          ))
        return result

      for row_offset, row in enumerate(data_rows):
        row_num = header_row_idx + 2 + row_offset
        question = self._get_cell(row, row_num, q_col, merged_map)
        answer = self._get_cell(row, row_num, a_col, merged_map)

        if not question and not answer:
          continue
        if is_section_heading(question, answer):
          section_heading = question
          continue
        if not question or not answer:
          continue

        result.qa_pairs.append(QAPairDraft(
          question=question,
          answer=answer,
          source_sheet=sheet_name,
          source_row_number=row_num,
          section_heading=section_heading,
        ))
        text_lines.append(f"Q: {question}\nA: {answer}")
    else:
      return self._two_column_fallback(rows, sheet_name, merged_map)

    if text_lines:
      result.full_text = f"=== {sheet_name} ===\n" + "\n\n".join(text_lines)
    return result

  def _two_column_fallback(
    self, rows: list[tuple], sheet_name: str, merged_map: dict
  ) -> ExtractionResult:
    result = ExtractionResult(method="excel_two_column")
    section_heading: str | None = None
    paired = 0
    text_lines: list[str] = []

    for row_num, row in enumerate(rows, start=1):
      question = self._get_cell(row, row_num, 0, merged_map)
      answer = self._get_cell(row, row_num, 1, merged_map)

      if question and answer:
        paired += 1
      if paired < 3 and not (question and answer):
        continue

      if not question and not answer:
        continue
      if is_section_heading(question, answer):
        section_heading = question
        continue
      if not question or not answer:
        continue

      result.qa_pairs.append(QAPairDraft(
        question=question,
        answer=answer,
        source_sheet=sheet_name,
        source_row_number=row_num,
        section_heading=section_heading,
      ))
      text_lines.append(f"Q: {question}\nA: {answer}")

    if text_lines:
      result.full_text = f"=== {sheet_name} ===\n" + "\n\n".join(text_lines)
    if not result.qa_pairs:
      result.warnings.append(f"Sheet '{sheet_name}': no Q&A pairs extracted")
    return result
