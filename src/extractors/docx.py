"""DOCX document extractor."""

import io
import logging

from docx import Document as DocxDocument

from src.extractors.base import ChunkDraft, ExtractionResult

logger = logging.getLogger(__name__)


class DocxExtractor:
    method = "docx"

    def extract(self, file_data: bytes, extension: str = ".docx") -> ExtractionResult:
        result = ExtractionResult(method="docx")
        try:
            doc = DocxDocument(io.BytesIO(file_data))
            current_heading: str | None = None
            chunk_parts: list[str] = []
            chunk_index = 0
            all_text: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style_name = para.style.name if para.style else ""
                if style_name.startswith("Heading"):
                    if chunk_parts:
                        content = "\n".join(chunk_parts)
                        result.chunks.append(ChunkDraft(
                            chunk_index=chunk_index,
                            content=content,
                            section_heading=current_heading,
                            content_type="text",
                        ))
                        chunk_index += 1
                        chunk_parts = []
                    current_heading = text
                    all_text.append(f"## {text}")
                else:
                    chunk_parts.append(text)
                    all_text.append(text)

            for table in doc.tables:
                table_rows: list[str] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    table_text = "\n".join(table_rows)
                    result.chunks.append(ChunkDraft(
                        chunk_index=chunk_index,
                        content=table_text,
                        section_heading=current_heading,
                        content_type="table",
                    ))
                    chunk_index += 1
                    all_text.append(table_text)

            if chunk_parts:
                content = "\n".join(chunk_parts)
                result.chunks.append(ChunkDraft(
                    chunk_index=chunk_index,
                    content=content,
                    section_heading=current_heading,
                    content_type="text",
                ))

            result.full_text = "\n\n".join(all_text)
            result.status = "completed" if result.full_text else "partial"
            result.metadata = {"chunk_count": len(result.chunks)}
        except Exception as exc:
            logger.exception("DOCX extraction failed")
            result.status = "failed"
            result.warnings.append(f"DOCX extraction error: {exc}")
        return result
