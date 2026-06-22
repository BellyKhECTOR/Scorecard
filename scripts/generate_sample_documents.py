#!/usr/bin/env python3
"""Generate the six acceptance corpus sample documents."""

import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from src.config import get_settings

SAMPLES = {
    "EPPF Asset Manager Questionnaire Domestic Equity_2025.pdf": "pdf_blank",
    "EPPF Asset Manager Questionnaire Domestic Equity_2025.docx": "docx_blank",
    "EPPF Asset Manager Questionnaire Domestic Equity_2025 Final.pdf": "pdf_final",
    "Investment Manager DD Questionnaire - 19052026.xlsx": "investment_dd",
    "ESG Questionnaire.xlsx": "esg",
    "Completed_HIM_RI_Questionnaire.xlsx": "him_ri",
}


def _write_pdf(path: Path, lines: list[str]):
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=(612, 792))
    y = 750
    for line in lines:
        c.drawString(72, y, line[:90])
        y -= 18
        if y < 72:
            c.showPage()
            y = 750
    c.save()
    path.write_bytes(buffer.getvalue())


def make_blank_pdf(path: Path):
    _write_pdf(path, [
        "EPPF Asset Manager Questionnaire - Domestic Equity",
        "Blank Template - 2025",
        "",
        "Section 1: Firm Information",
        "Question: Company name?",
        "Response:",
        "",
        "Question: AUM?",
        "Response:",
    ])


def make_final_pdf(path: Path):
    _write_pdf(path, [
        "EPPF Asset Manager Questionnaire - Domestic Equity - FINAL",
        "Mazi Asset Management (Pty) Ltd",
        "",
        "1. What is your firm's AUM?",
        "Mazi manages approximately R15 billion in assets under management.",
        "",
        "2. Describe your investment process.",
        "Mazi follows a disciplined bottom-up fundamental research process.",
        "",
        "3. ESG integration approach?",
        "ESG factors are integrated into our investment analysis and decision-making.",
        "",
        "4. Proxy voting policy?",
        "We vote all proxies in accordance with our published policy.",
    ])


def make_docx(path: Path):
    doc = DocxDocument()
    doc.add_heading("EPPF Asset Manager Questionnaire - Domestic Equity", level=1)
    doc.add_paragraph("Section: Firm Overview")
    doc.add_paragraph("Company name: Mazi Asset Management (Pty) Ltd")
    doc.add_paragraph("This is a blank questionnaire template for domestic equity mandate.")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Question"
    table.rows[0].cells[1].text = "Response"
    table.rows[1].cells[0].text = "AUM?"
    table.rows[1].cells[1].text = ""
    table.rows[2].cells[0].text = "Team size?"
    table.rows[2].cells[1].text = ""
    doc.save(path)


def make_esg_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ESG"
    ws.append(["Question", "Response"])
    ws.append(["Does the firm have a formal ESG policy?", "Yes, adopted in 2020 and reviewed annually"])
    ws.append(["How is climate risk integrated?", "Scenario analysis conducted annually"])
    ws.append(["GOVERNANCE", ""])
    ws.append(["Board oversight of ESG?", "Quarterly board review of ESG matters"])
    wb.save(path)


def make_him_ri_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RI Questionnaire"
    ws.append(["Question", "Indicator", "Confirmation", "Response"])
    ws.append(["Proxy voting policy in place?", "Y/N", "Y", "We vote all proxies in line with our policy"])
    ws.append(["Signatory to stewardship code?", "Y/N", "Y", "Yes, signatory since 2019"])
    ws.append(["Active ownership approach?", "Y/N", "Y", "Engage with companies on material ESG issues"])
    ws.append(["Stewardship reporting?", "Y/N", "Y", "Annual stewardship report published"])
    wb.save(path)


def make_investment_dd_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Firm Overview"
    ws1.append(["Company legal name?", "Mazi Asset Management (Pty) Ltd"])
    ws1.append(["AUM?", "R15 billion"])
    ws1.append(["Year founded?", "2010"])
    ws1.append(["FIRM GOVERNANCE", ""])
    ws1.append(["Board composition?", "Independent chair and majority independent directors"])

    ws2 = wb.create_sheet("Investment Process")
    ws2.append(["Question", "Indicator", "Annexure", "Response"])
    ws2.append(["Investment philosophy?", "", "", "Bottom-up fundamental value investing"])
    ws2.append(["Risk management framework?", "", "", "Three lines of defence model"])

    ws3 = wb.create_sheet("Metadata")
    ws3.sheet_state = "hidden"
    ws3.append(["internal", "do not extract"])
    wb.save(path)


def main() -> int:
    settings = get_settings()
    out_dir = settings.sample_documents_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"Generating sample documents in: {out_dir}")

    generators = {
        "pdf_blank": make_blank_pdf,
        "docx_blank": make_docx,
        "pdf_final": make_final_pdf,
        "esg": make_esg_xlsx,
        "him_ri": make_him_ri_xlsx,
        "investment_dd": make_investment_dd_xlsx,
    }

    for filename, kind in SAMPLES.items():
        path = out_dir / filename
        generators[kind](path)
        print(f"  Created: {filename}")

    print(f"\nSUCCESS: {len(SAMPLES)} sample documents ready.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
