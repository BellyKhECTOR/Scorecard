"""Create synthetic test fixtures for extractors."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfWriter


def create_esg_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ESG"
    ws.append(["Question", "Response"])
    ws.append(["Does the firm have an ESG policy?", "Yes, adopted in 2020"])
    ws.append(["How is climate risk integrated?", "Scenario analysis annually"])
    ws.append(["GOVERNANCE", ""])
    ws.append(["Board ESG oversight?", "Quarterly board review"])
    wb.save(path)


def create_him_ri_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RI Questionnaire"
    ws.append(["Question", "Indicator", "Confirmation", "Response"])
    ws.append(["Proxy voting policy?", "Y/N", "Y", "We vote all proxies"])
    ws.append(["Stewardship code signatory?", "Y/N", "Y", "Yes since 2019"])
    ws.append(["Active ownership approach?", "Y/N", "Y", "Engage on ESG issues"])
    wb.save(path)


def create_two_column_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Company name?", "Mazi Asset Management"])
    ws.append(["AUM?", "R15 billion"])
    ws.append(["Founded?", "2010"])
    wb.save(path)


def create_simple_docx(path: Path):
    doc = DocxDocument()
    doc.add_heading("Investment Process", level=1)
    doc.add_paragraph("Mazi follows a disciplined bottom-up research process.")
    doc.add_paragraph("The team consists of 12 investment professionals.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "AUM"
    table.rows[1].cells[1].text = "R15bn"
    doc.save(path)


def create_simple_pdf(path: Path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as f:
        writer.write(f)


def main():
    fixtures_dir = Path(__file__).resolve().parent
    fixtures_dir.mkdir(exist_ok=True)
    create_esg_xlsx(fixtures_dir / "esg_questionnaire.xlsx")
    create_him_ri_xlsx(fixtures_dir / "him_ri_questionnaire.xlsx")
    create_two_column_xlsx(fixtures_dir / "two_column_ddq.xlsx")
    create_simple_docx(fixtures_dir / "sample.docx")
    create_simple_pdf(fixtures_dir / "sample.pdf")
    print(f"Fixtures created in {fixtures_dir}")


if __name__ == "__main__":
    main()
